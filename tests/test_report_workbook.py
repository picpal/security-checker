"""5 시트 빌더(spec §5.5) — 열 계약, 사실/추정 분리, 비밀값 미노출, 결정성, 골든."""
import json
from pathlib import Path

from secscan.output.json_io import from_json
from secscan.report.deppath import DepGraph, deppaths_for
from secscan.report.derive import derive_all
from secscan.report.kb import entry_for
from secscan.report.models import Interpretation
from secscan.report.workbook import (ACTION_HEADER, QUESTION_HEADER, RESULT_HEADER, SHEET_ACTIONS, SHEET_QUESTIONS,
                                     SHEET_RESULT, SHEET_RULES, SHEET_SUMMARY, build_report_sheets, result_notes, write_report)

FIX = Path("fixtures/report/work-note-findings.json")
GOLDEN = Path("tests/golden/report-work-note-sheets.json")
FS = from_json(FIX.read_text(encoding="utf-8"))
META = json.loads(FIX.read_text(encoding="utf-8"))["meta"]
KBS = {f.id: entry_for(f) for f in FS}
DER = derive_all(FS, KBS, deppaths_for(FS, DepGraph.from_path("fixtures/report/work-note-bom.cdx.json")))


def _sheets(interps=None, interp_meta=None):
    return build_report_sheets(FS, META, KBS, DER, interps, interp_meta)


def test_sheet_names_and_headers():
    s = _sheets()
    assert list(s) == [SHEET_SUMMARY, SHEET_ACTIONS, SHEET_QUESTIONS, SHEET_RULES, SHEET_RESULT]
    assert s[SHEET_ACTIONS][0] == ACTION_HEADER and len(ACTION_HEADER) == 17
    assert s[SHEET_QUESTIONS][0] == QUESTION_HEADER and len(QUESTION_HEADER) == 9
    assert s[SHEET_RESULT][0] == RESULT_HEADER and len(RESULT_HEADER) == 9
    assert ACTION_HEADER[5] == "의존 경로" and ACTION_HEADER[11] == "같이 해결"


def test_action_rows_sorted_and_first_row_is_hardcoded_credential():
    rows = _sheets()[SHEET_ACTIONS][1]
    assert len(rows) == 15 and rows[0][0] == 1 and rows[0][1] == "P1" and rows[0][2] == "356e4bbab00c"
    assert rows[0][3] == "소스코드에 박힌 자격증명" and rows[0][10] == "사람 확인 후"
    assert rows[0][12] == "SAST · 룰 신뢰도 high → tier actionable → 조치대상"
    assert rows[0][15] == "KISA SF-06 하드코드된 중요정보 · PCI-DSS 6.2.4 — access control"


def test_question_rows_only_for_needs_question_and_interpretation_columns_empty_without_file():
    rows = _sheets()[SHEET_QUESTIONS][1]
    assert len(rows) == 15  # 사람 확인 후 11 + evidence 없는 강등 SCA 4
    for r in rows:
        assert r[5] == "" and r[6] == "" and r[7] == ""  # 추정·확인 방법·인용 비어 있음
        assert r[8].startswith("① 수정")


def test_interpretations_fill_guess_columns_and_summary_reports_status():
    interps = {"356e4bbab00c": Interpretation("상수 PASSWORD 가 KNOWN_KEYS 집합에 들어감.", "88행 리터럴 값 확인.",
                                             ("backend/src/main/java/com/worknote/admin/BreakGlassFile.java:87-89",))}
    s = _sheets(interps, {"status": "ok", "rejected": [{"id": "108d8ee4651c", "reason": "no_cite"}]})
    row = next(r for r in s[SHEET_QUESTIONS][1] if r[1] == "356e4bbab00c")
    assert row[5].startswith("상수 PASSWORD") and row[6] == "88행 리터럴 값 확인." and "87-89" in row[7]
    summary = {r[0]: (r[1], r[2]) for r in s[SHEET_SUMMARY][1]}
    assert summary["해석 상태"][0] == "ok" and "108d8ee4651c: no_cite" in summary["해석 상태"][1]


def test_summary_conclusion_and_counts_and_unregistered_rules():
    summary = {r[0]: (r[1], r[2]) for r in _sheets()[SHEET_SUMMARY][1]}
    assert summary["결론"][0].startswith("총 15건. 즉시 조치 대상 8건(자동 수정 0, 사람 확인 후 8).")
    assert summary["즉시 조치 대상(P1·P2)"][0] == 8 and summary["강등(P4)"][0] == 4 and summary["제외"][0] == 23
    assert summary["스캐너 trivy"][0] == "ok"
    assert summary["지식베이스 미등록 룰"][0] == "없음"  # 10건 등록으로 work-note 는 전부 커버


def test_result_sheet_is_pure_table_and_notes_live_outside():
    rows = _sheets()[SHEET_RESULT][1]
    acts = _sheets()[SHEET_ACTIONS][1]
    assert [r[0] for r in rows] == [r[2] for r in acts]  # 안내 행 없이 finding 행만, 취약점 목록 순서
    assert all(r[1:6] == ["", "", "", "", ""] for r in rows)  # 처리 결과~커밋은 LLM 이 채움
    assert [r[6:] for r in rows] == [[a[3], a[10], a[4]] for a in acts]  # 문제·담당·위치 참고 열
    notes = result_notes("work-note")
    assert [n[0] for n in notes] == ["이 시트는", "고쳐도 되는 것", "하지 말 것", "끝났는지 확인하는 법", "처리 결과에 쓸 수 있는 값", "검증은 누가 하나"]
    assert "secscan scan --target work-note" in notes[3][1]


def test_rules_sheet_is_fixed_table():
    header, rows = _sheets()[SHEET_RULES]
    assert header == ["구분", "값·용어", "뜻(일반어)", "그래서 어떻게 하나"]
    groups = [r[0] for r in rows]
    assert groups[:5] == ["우선순위"] * 5 and {"담당", "판정", "심각도", "조치 유형", "같이 해결", "사실/추정", "주의"} <= set(groups)
    assert all(len(r) == 4 and all(isinstance(c, str) and c for c in r) for r in rows)


def test_no_secret_values_and_all_cells_sanitized():
    s = _sheets()
    blob = json.dumps({k: v[1] for k, v in s.items()}, ensure_ascii=False)
    assert "pend-" not in blob and "ghp_" not in blob
    for _, rows in s.values():
        for r in rows:
            for c in r:
                assert c is None or isinstance(c, (int, float)) or (isinstance(c, str) and not c.lstrip().startswith(("=", "+", "-", "@")))


def test_deterministic_regardless_of_input_order():
    rev = list(reversed(FS))
    dp = deppaths_for(FS, DepGraph.from_path("fixtures/report/work-note-bom.cdx.json"))
    assert build_report_sheets(rev, META, {f.id: KBS[f.id] for f in rev}, derive_all(rev, KBS, dp)) == _sheets()


def test_matches_golden():
    got = {k: {"header": h, "rows": r} for k, (h, r) in _sheets().items()}
    assert got == json.loads(GOLDEN.read_text(encoding="utf-8"))


def test_write_report_xlsx_or_csv(tmp_path):
    paths, warn = write_report(_sheets(), tmp_path, created="2026-09-06", prefer_xlsx=False)
    assert warn and (tmp_path / "report-xlsx" / "1_취약점 목록.csv").exists()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return
    paths, warn = write_report(_sheets(), tmp_path, created="2026-09-06", prefer_xlsx=True, notes=result_notes("work-note"))
    assert warn is None and paths == [tmp_path / "report.xlsx"]
    from openpyxl import load_workbook
    wb = load_workbook(paths[0])
    rs = wb[SHEET_RESULT]
    assert rs["A1"].value == "이 시트는" and rs["A7"].value is None and rs["A8"].value == "ID"  # 안내 6행 + 빈 행 + 헤더
    assert rs["A9"].value == _sheets()[SHEET_RESULT][1][0][0] and rs.freeze_panes == "A9"
    assert wb.sheetnames == [SHEET_SUMMARY, SHEET_ACTIONS, SHEET_QUESTIONS, SHEET_RULES, SHEET_RESULT]
    assert wb.properties.modified == wb.properties.created
    from openpyxl.utils import get_column_letter
    from secscan.report.workbook import MAX_WIDTH, MIN_WIDTH, fit_width
    ws = wb[SHEET_ACTIONS]
    widths = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, len(ACTION_HEADER) + 1)]
    assert all(MIN_WIDTH <= w <= MAX_WIDTH for w in widths) and widths[0] < widths[8]  # No. 열은 좁고 '해야 할 일' 열은 넓다
    assert fit_width(["ab", "한글한글"]) == 10.0 and fit_width([]) == MIN_WIDTH and fit_width(["x" * 200]) == MAX_WIDTH
    assert ws.freeze_panes == "D2" and ws["A1"].font.bold
