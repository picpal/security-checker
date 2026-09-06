"""xlsx projection (spec §8) — 보안팀 12열 Deduped_CVEs 와 1:1 + secscan 열.

이 모듈의 순수부(build_sheets)는 stdlib 만 쓴다. openpyxl 은 write_xlsx 에서만 지연 import 하고,
없으면 CSV 번들로 폴백한다(코어 의존성 0 유지). 출력은 disposition/tier 필드만 읽는다(재계산 금지).
비밀값(secret 원문·fingerprint)은 어느 시트에도 쓰지 않는다.
"""
from __future__ import annotations

import re
from typing import Iterable

from .. import __version__
from ..models import Finding, severity_rank
from ._status import status_rows

Sheet = tuple[list[str], list[list]]

MAX_ROWS = 1_048_576  # Excel 시트 행 상한(헤더 포함)
MAX_CHARS = 32_767  # Excel 셀 문자 상한
_TRUNC = "…[truncated]"
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_FORMULA_LEAD = ("=", "+", "-", "@")

TEAM_HEADER = ("No.", "Vulnerability ID", "Severity", "Max CVSS Score", "CVSS Source", "Affected Packages",
               "Installed Versions", "Fixed Versions", "Targets", "Finding Count", "Title", "URL")
SECSCAN_EXTRA_HEADER = ("Published", "Reachability", "Reachability Evidence", "Consensus Tools", "Consensus Score",
                        "Disposition", "KISA", "PCI", "Suppression", "Finding ID")
FINDINGS_HEADER = ["Finding ID", "Category", "Severity", "Rule/Advisory", "Title", "Package", "Installed Version",
                   "Fixed Versions", "Target", "Tool", "File", "Line", "CWE", "Reachability", "Disposition", "Tier", "Suppressed"]
SAST_HEADER = ["Finding ID", "Rule", "Severity", "Confidence", "Tier", "File", "Line", "CWE", "KISA", "PCI", "Suppressed", "Disposition"]
SECRET_HEADER = ["Finding ID", "Rule", "File", "Line", "Verified", "Severity", "Suppressed", "Disposition"]


def sanitize_cell(v):
    """formula injection 방어 + XML 불법 문자 제거 + 길이 절단. 숫자/None 은 그대로, bool 은 문자열."""
    if v is None or isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    s = str(v)
    s = _ILLEGAL_XML.sub("", s)
    if s.lstrip(" \t\r\n").startswith(_FORMULA_LEAD):
        s = "'" + s
    if len(s) > MAX_CHARS:
        s = s[: MAX_CHARS - len(_TRUNC)] + _TRUNC
    return s


def _join(values: Iterable[str]) -> str:
    return ", ".join(sorted({v for v in values if v}))


def _kisa(f: Finding) -> str:
    return _join(w.name for w in f.compliance.kisa) if f.compliance else ""


def _pci(f: Finding) -> str:
    return _join(f.compliance.pci) if f.compliance else ""


def _sup(f: Finding) -> str:
    return f"{f.suppression.state}: {f.suppression.reason}" if f.suppression else ""


def _occ(f: Finding):
    if f.occurrences:
        return list(f.occurrences)
    # occurrence 가 없는 SCA(옛 증거) 는 component 기준 1건으로 본다
    from ..models import Occurrence
    if f.component:
        return [Occurrence(f.tool, f.source or "", f.component.package, f.component.version)]
    return []


def _max_cvss(f: Finding):
    cv = [c for c in (f.advisory.cvss if f.advisory else ()) if c.score is not None]
    if not cv:
        return None, ""
    best = sorted(cv, key=lambda c: (-c.score, c.source))[0]
    return best.score, best.source


def _rows_findings(findings: list[Finding]) -> list[list]:
    rows = []
    for f in findings:
        base = [f.id, f.category, f.severity, f.rule_id, f.title]
        tail = [_join(f.cwe), f.reachability.status, f.disposition or "", f.tier or "", _sup(f)]
        if f.category == "sca":
            for o in _occ(f):
                rows.append(base + [o.package, o.version, _join(f.advisory.fixed_versions if f.advisory else ()), o.target,
                                    o.tool, "", None] + tail)
        else:
            rows.append(base + ["", "", "", "", f.tool, f.location.file if f.location else "",
                                f.location.start_line if f.location else None] + tail)
    return rows


def _rows_deduped(findings: list[Finding]) -> list[list]:
    groups: dict[str, list[Finding]] = {}
    for f in findings:
        if f.category == "sca" and f.advisory:
            groups.setdefault(f.advisory.id, []).append(f)
    rows = []
    for n, (adv_id, fs) in enumerate(sorted(groups.items()), start=1):
        occ = [o for f in fs for o in _occ(f)]
        sev = max((f.severity for f in fs), key=severity_rank)
        scores = [(_max_cvss(f)) for f in fs]
        scored = [s for s in scores if s[0] is not None]
        score, source = (sorted(scored, key=lambda s: (-s[0], s[1]))[0] if scored else (None, ""))
        head = fs[0]
        rows.append([
            n, adv_id, sev, score, source, _join(o.package for o in occ), _join(o.version for o in occ),
            _join(v for f in fs for v in f.advisory.fixed_versions), _join(o.target for o in occ), len(occ) or len(fs),
            head.title, head.references[0] if head.references else "",
            head.advisory.published or "", _join(f.reachability.status for f in fs),
            _join(f.reachability.evidence or "" for f in fs), _join(t for f in fs if f.consensus for t in f.consensus.tools),
            max((f.consensus.score for f in fs if f.consensus), default=0), _join(f.disposition or "" for f in fs),
            _join(_kisa(f) for f in fs), _join(_pci(f) for f in fs), _join(_sup(f) for f in fs), _join(f.id for f in fs),
        ])
    return rows


def _rows_targets(findings: list[Finding]) -> list[list]:
    counts: dict[str, int] = {}
    for f in findings:
        for o in _occ(f) if f.category == "sca" else []:
            counts[o.target] = counts.get(o.target, 0) + 1
    return [[t, n] for t, n in sorted(counts.items())]


def _rows_sast(findings: list[Finding]) -> list[list]:
    return [[f.id, f.rule_id, f.severity, f.confidence, f.tier or "", f.location.file if f.location else "",
             f.location.start_line if f.location else None, _join(f.cwe), _kisa(f), _pci(f), _sup(f), f.disposition or ""]
            for f in findings if f.category == "sast"]


def _rows_secret(findings: list[Finding]) -> list[list]:
    # 비밀값·fingerprint(snippet_ref) 를 쓰지 않는다.
    return [[f.id, f.rule_id, f.location.file if f.location else "", f.location.start_line if f.location else None,
             "" if f.verified is None else str(f.verified), f.severity, _sup(f), f.disposition or ""]
            for f in findings if f.category == "secret"]


def _rows_summary(findings: list[Finding], meta: dict) -> list[list]:
    from collections import Counter
    rows: list[list] = []
    for (cat, sev), n in sorted(Counter((f.category, f.severity) for f in findings).items()):
        rows.append(["count", f"{cat}/{sev}", n])
    for d, n in sorted(Counter(f.disposition or "undecided" for f in findings).items()):
        rows.append(["disposition", d, n])
    for t, n in sorted(Counter(f.tier for f in findings if f.tier).items()):
        rows.append(["tier", t, n])
    for s, n in sorted(Counter(f.reachability.status for f in findings if f.category == "sca").items()):
        rows.append(["reachability", s, n])
    rows += [["profile", "", meta.get("profile", "")], ["target", "", meta.get("target", "")],
             ["scanned_at", "", meta.get("run_date", "")], ["secscan_version", "", __version__]]
    return rows


def _rows_meta(meta: dict) -> list[list]:
    rows: list[list] = []
    statuses = sorted(status_rows(meta), key=lambda s: s["tool"])
    for s in statuses:
        for k in ("status", "tool_version", "duration_s", "message"):
            rows.append([f"scanner.{s['tool']}.{k}", s[k] if s[k] is not None else ""])
    rows += [["profile", meta.get("profile", "")], ["commit", meta.get("commit", meta.get("snapshot", ""))],
             ["partial_failures", _join(s["tool"] for s in statuses if s["status"] != "ok")],
             ["reachability.ran", str(meta.get("reachability", {}).get("ran", ""))],
             ["reachability.reason", meta.get("reachability", {}).get("reason", "")],
             ["secret_policy", meta.get("secret_policy", "")], ["excluded_count", meta.get("excluded_count", "")]]
    return rows


def _san(rows: list[list]) -> list[list]:
    return [[sanitize_cell(c) for c in r] for r in rows]


def build_sheets(findings: list[Finding], meta: dict) -> dict[str, Sheet]:
    """순수·결정적. 입력 순서와 무관하게 같은 시트를 만든다(각 시트가 자체 정렬)."""
    from .order import sort_findings
    fs = sort_findings(findings)
    return {
        "Summary": (["Metric", "Key", "Value"], _san(_rows_summary(fs, meta))),
        "Findings": (FINDINGS_HEADER, _san(_rows_findings(fs))),
        "Deduped_CVEs": (list(TEAM_HEADER + SECSCAN_EXTRA_HEADER), _san(_rows_deduped(fs))),
        "Targets": (["Target", "Findings"], _san(_rows_targets(fs))),
        "SAST": (SAST_HEADER, _san(_rows_sast(fs))),
        "Secret": (SECRET_HEADER, _san(_rows_secret(fs))),
        "Meta": (["Key", "Value"], _san(_rows_meta(meta))),
    }


# --- 라이터부: 시트 분할·CSV 폴백·xlsx(openpyxl 지연 import) ---

import csv
import zipfile
from datetime import datetime
from pathlib import Path

_MODIFIED_RE = re.compile(rb'(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)')


def split_sheet(name: str, sheet: Sheet, max_rows: int = MAX_ROWS) -> list[tuple[str, Sheet]]:
    header, rows = sheet
    per = max(1, max_rows - 1)  # 헤더 1행 제외
    if len(rows) <= per:
        return [(name, sheet)]
    return [(name if i == 0 else f"{name}_{i + 1}", (header, rows[i * per:(i + 1) * per]))
            for i in range((len(rows) + per - 1) // per)]


def _split_all(sheets: dict[str, Sheet], max_rows: int = MAX_ROWS) -> list[tuple[str, Sheet]]:
    out: list[tuple[str, Sheet]] = []
    for name, sheet in sheets.items():
        out += split_sheet(name, sheet, max_rows)
    return out


def write_csv_bundle(sheets: dict[str, Sheet], out_dir, *, max_rows: int = MAX_ROWS) -> list[Path]:
    """openpyxl 부재 시 폴백. 같은 sanitizer 를 이미 거친 셀을 그대로 쓴다."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    written = []
    for name, (header, rows) in _split_all(sheets, max_rows):
        p = out / f"{name}.csv"
        with p.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
            w.writerow(header)
            w.writerows(["" if c is None else c for c in r] for r in rows)
        written.append(p)
    return written


def _pin_xlsx_modified(path: Path, ts: datetime) -> None:
    """M11(최종 리뷰) — openpyxl 은 `wb.save()` 시점에 `docProps/core.xml` 의 `dcterms:modified`
    를 저장 시각(now)으로 덮어쓴다 — `wb.properties.modified` 를 미리 고정해도 무시된다. 실행마다
    값이 달라지면 spec §8 "workbook.properties.created/modified 를 스캔 일시로 고정"을 어긴다.
    save() 가 끝난 뒤 zip 을 열어 그 한 요소만 다시 쓴다 — 다른 멤버는 그대로 보존한다."""
    iso = ts.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    with zipfile.ZipFile(path, "r") as zin:
        infos = zin.infolist()
        payloads = {info.filename: zin.read(info.filename) for info in infos}
    core = payloads.get("docProps/core.xml")
    if core is None:
        return
    payloads["docProps/core.xml"] = _MODIFIED_RE.sub(lambda m: m.group(1) + iso + m.group(2), core)
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, payloads[info.filename])
    tmp.replace(path)


def write_xlsx(sheets: dict[str, Sheet], path, *, created: str, max_rows: int = MAX_ROWS) -> Path:
    """openpyxl 필요(`pip install secscan[xlsx]`). 문자열 셀은 data_type 's' 로 강제 — '=' 시작 문자열이 수식이 되지 않게."""
    from openpyxl import Workbook  # 지연 import: 코어 의존성 0 유지
    from openpyxl.cell.cell import TYPE_STRING

    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in _split_all(sheets, max_rows):
        ws = wb.create_sheet(title=name[:31])
        ws.append(header)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = TYPE_STRING
    ts = datetime.fromisoformat(created) if created else datetime(1970, 1, 1)
    wb.properties.created = ts
    wb.properties.modified = ts
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    _pin_xlsx_modified(p, ts)
    return p


def _openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def write_workbook(findings: list[Finding], meta: dict, out_dir, *, prefer_xlsx: bool | None = None,
                   max_rows: int = MAX_ROWS) -> tuple[list[Path], str | None]:
    """xlsx 가 가능하면 `<out>/findings.xlsx`, 아니면 `<out>/findings-xlsx/<sheet>.csv` + 경고 1줄."""
    sheets = build_sheets(findings, meta)
    out = Path(out_dir)
    use_xlsx = _openpyxl_available() if prefer_xlsx is None else prefer_xlsx
    if use_xlsx:
        return [write_xlsx(sheets, out / "findings.xlsx", created=str(meta.get("run_date", "")), max_rows=max_rows)], None
    paths = write_csv_bundle(sheets, out / "findings-xlsx", max_rows=max_rows)
    return paths, "openpyxl 미설치 — 시트별 CSV 번들로 대체(`pip install secscan[xlsx]` 로 xlsx 출력)"
