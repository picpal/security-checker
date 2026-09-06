"""결과 반환(LLM) ↔ 재스캔 findings.json 대조(spec §5.6). fixed 주장은 id 부재로만 인정한다."""
from __future__ import annotations

import json
from pathlib import Path

from ..output.json_io import from_json
from .models import RESULT_STATUSES, ResultRow
from .paths import relativize
from .workbook import SHEET_RESULT

RESULT_CONTEXT = "secscan-report-result/v1"
_FIELDS = ("id", "status", "changed_files", "verification", "reason", "commit")
_XLSX_COLS = ("id", "status", "changed_files", "verification", "reason", "commit")  # 4_결과반환 열 순서


def _row(d: dict) -> ResultRow:
    d = {k: ("" if d.get(k) is None else str(d.get(k))) for k in _FIELDS}
    return ResultRow(**d)


def load_result(path) -> list[ResultRow]:
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError("xlsx 결과를 읽으려면 openpyxl 이 필요합니다(`pip install secscan[xlsx]`)") from e
        ws = load_workbook(p, read_only=True)[SHEET_RESULT]
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not r or r[0] is None or str(r[0]).startswith("#"):
                continue
            rows.append(_row(dict(zip(_XLSX_COLS, list(r) + [""] * 6))))
        return rows
    data = json.loads(p.read_text(encoding="utf-8"))
    items = data.get("rows", []) if isinstance(data, dict) else data
    return [_row(d) for d in items]


def rescan_ids(findings_path, target: str | None = None) -> set[str]:
    """재스캔 findings.json 을 --target(없으면 meta.target) 기준으로 상대화한 뒤 id 집합을 낸다.

    보고서·report-request.json·4_결과반환 의 ID 는 상대화 공간이므로, 여기서 같은 공간으로
    맞추지 않으면 fixed 주장이 전부 거짓으로 '일치' 판정된다(spec §12(f))."""
    text = Path(findings_path).read_text(encoding="utf-8")
    findings = from_json(text)
    meta_target = json.loads(text).get("meta", {}).get("target")
    findings = relativize(findings, target if target is not None else meta_target)
    return {f.id for f in findings}


def check(rows: list[ResultRow], present_ids: set[str]) -> list[dict]:
    out = []
    for r in rows:
        is_present = r.id in present_ids
        if r.status not in RESULT_STATUSES:
            verdict = "불일치(status 값 오류)"
        elif r.status == "fixed":
            verdict = "불일치" if is_present else "일치"
        else:
            verdict = "미검증" if is_present else "미검증(재스캔에서 사라짐)"
        out.append({"id": r.id, "status": r.status, "present": is_present, "verdict": verdict})
    return out


def has_mismatch(results: list[dict]) -> bool:
    return any(r["verdict"].startswith("불일치") for r in results)


def render(results: list[dict]) -> str:
    lines = ["| ID | status | 재스캔 존재 | 판정 |", "|---|---|---|---|"]
    lines += [f"| {r['id']} | {r['status']} | {'존재' if r['present'] else '부재'} | {r['verdict']} |" for r in results]
    return "\n".join(lines)
